from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK

OUT = "/Users/preetinagpal/Downloads/ResumeBuilder/CareerSense_SEO_Next_Steps.docx"
NAVY = "0B2848"
BLUE = "123B6D"
GOLD = "A87922"
INK = "071426"
MUTED = "60738A"
PALE = "EEF3F8"
LIGHT = "F7F9FB"
BORDER = "D7E0EA"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1)
section.bottom_margin = Inches(0.85)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

def set_font(run, size=11, color=INK, bold=False, italic=False, name="Aptos"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18

for name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 9),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 11.5, NAVY, 10, 5),
]:
    st = styles[name]
    st.font.name = "Aptos Display"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for style_name in ["List Bullet", "List Number"]:
    st = styles[style_name]
    st.font.name = "Aptos"
    st.font.size = Pt(10.5)
    st.paragraph_format.left_indent = Inches(0.375)
    st.paragraph_format.first_line_indent = Inches(-0.188)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.18

if "Code Block" not in styles:
    code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code Block"]
code_style.font.name = "Consolas"
code_style._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
code_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor.from_string(NAVY)
code_style.paragraph_format.left_indent = Inches(0.18)
code_style.paragraph_format.right_indent = Inches(0.18)
code_style.paragraph_format.space_before = Pt(4)
code_style.paragraph_format.space_after = Pt(8)
code_style.paragraph_format.line_spacing = 1.05

def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)

def add_para(text="", bold_lead=None, style=None, color=INK, size=10.5, after=6, keep=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = keep
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, size=size, color=color, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r, size=size, color=color)
    else:
        r = p.add_run(text)
        set_font(r, size=size, color=color)
    return p

def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_font(p.add_run(item), size=10.5)

def add_steps(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        set_font(p.add_run(item), size=10.5)

def add_code(lines):
    p = doc.add_paragraph(style="Code Block")
    shade_paragraph(p, PALE)
    for idx, line in enumerate(lines):
        r = p.add_run(line)
        set_font(r, size=9, color=NAVY, name="Consolas")
        if idx < len(lines) - 1:
            r.add_break()

def add_callout(label, text, fill="FFF8E9"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.16)
    p.paragraph_format.right_indent = Inches(0.16)
    shade_paragraph(p, fill)
    r = p.add_run(f"{label}: ")
    set_font(r, size=10.5, color=GOLD, bold=True)
    r = p.add_run(text)
    set_font(r, size=10.5, color=INK)

def add_checklist(items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.05)
        p.paragraph_format.space_after = Pt(5)
        set_font(p.add_run("☐  "), size=11, color=GOLD, bold=True)
        set_font(p.add_run(item), size=10.5)

def add_table(headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = 0
    header = table.rows[0]
    prevent_row_split(header)
    for i, text in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(text), size=9.5, color="FFFFFF", bold=True)
    for row_idx, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for i, text in enumerate(values):
            cell = cells[i]
            cell.width = Inches(widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_idx % 2 == 1:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(str(text)), size=9.25, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

# Running furniture
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_font(hp.add_run("CAREERSENSE  /  SEO LAUNCH PLAYBOOK"), size=8.5, color=MUTED, bold=True)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(fp.add_run("CareerSense Resume Builder • Internal launch guide"), size=8, color=MUTED)

# Cover
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(115)
p.paragraph_format.space_after = Pt(18)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("SEO LAUNCH PLAYBOOK"), size=10, color=GOLD, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
set_font(p.add_run("CareerSense Resume Builder"), size=30, color=NAVY, bold=True, name="Aptos Display")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(24)
set_font(p.add_run("Deployment, Search Console, analytics, indexing, and ongoing optimization"), size=14, color=BLUE)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(30)
p.paragraph_format.space_after = Pt(8)
set_font(p.add_run("Prepared for the manual production launch"), size=10.5, color=MUTED, italic=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("October 2026"), size=11, color=NAVY, bold=True)
add_callout("Scope", "All code changes remain local. GitHub push, deployment, Search Console submission, and account-level configuration are performed manually by the site owner.", fill=PALE)
doc.add_page_break()

doc.add_heading("1. What is ready now", level=1)
add_para("The local implementation contains the full technical and content SEO foundation. The production build generates 62 prerendered public SEO pages and a 63-URL sitemap, including the homepage.")
add_table(
    ["Area", "Current state", "Launch requirement"],
    [
        ["SEO pages", "Role, student, institute, template, and guide pages", "Deploy generated output"],
        ["Technical SEO", "Canonicals, metadata, schema, robots, and sitemap", "Verify on production domain"],
        ["Analytics", "GA4 loader and conversion events implemented", "Provide actual measurement ID"],
        ["Search Console", "Priority URL list generated", "Verify property and submit sitemap"],
        ["GitHub", "No commit or push performed", "Review and push manually"],
    ],
    [1.35, 3.25, 1.9],
)
add_callout("Generated files", "After npm run build, use dist/sitemap.xml, dist/robots.txt, dist/seo-url-inventory.csv, and dist/search-console-priority-urls.txt.")

doc.add_heading("2. Pre-deployment configuration", level=1)
doc.add_heading("2.1 Confirm the production domain", level=2)
add_para("The implementation currently defaults to https://resume.careersenseai.com. If the final public domain differs, set VITE_SITE_URL before building. The value must include https:// and must not end with a slash.")
doc.add_heading("2.2 Create or identify the GA4 web stream", level=2)
add_steps([
    "Open Google Analytics and select the correct account and property.",
    "Create a Web data stream for the production domain, or open the existing stream.",
    "Copy the measurement ID in the format G-XXXXXXXXXX.",
    "Add the production values to the deployment environment; do not commit secrets or environment-specific values unnecessarily.",
])
add_code([
    "VITE_SITE_URL=https://resume.careersenseai.com",
    "VITE_GA_MEASUREMENT_ID=G-YOUR-ACTUAL-ID",
])
add_callout("Analytics behavior", "If VITE_GA_MEASUREMENT_ID is absent, the site works normally and analytics remains disabled. Never deploy the placeholder G-XXXXXXXXXX value.")

doc.add_heading("3. Review and push the local changes", level=1)
add_para("Perform this step manually, as requested. Review the complete working tree carefully because the repository also contains pre-existing template changes beyond the SEO work.")
add_steps([
    "Run git status --short and identify the intended files.",
    "Review git diff for tracked files and inspect every untracked file before staging.",
    "Remove local-only files such as .DS_Store from the staging set and add them to .gitignore if appropriate.",
    "Run the production build and acceptance checks before committing.",
    "Stage only the files intended for this release, create the commit, and push the selected branch manually.",
])
add_code([
    "git status --short",
    "git diff --check",
    "npm run build",
    "# Review files, then stage and commit manually",
])
add_callout("Do not use a broad git add command blindly", "The workspace includes unrelated or pre-existing changes. Confirm the release scope file by file before staging.", fill="FCEFD9")

doc.add_heading("4. Deploy to Vercel", level=1)
add_steps([
    "Open the Vercel project connected to the intended GitHub repository.",
    "Confirm the framework preset is Vite and the output directory is dist.",
    "Add VITE_SITE_URL and VITE_GA_MEASUREMENT_ID to the Production environment.",
    "Deploy the reviewed commit. If automatic Git deployments are enabled, monitor the deployment triggered by the push.",
    "Open the deployment logs and confirm the build reports 62 prerendered SEO pages and 63 sitemap URLs.",
    "Promote the deployment only after the preview checks pass, if your workflow uses preview deployments.",
])
add_para("The updated Vercel routing uses the filesystem first, then falls back to index.html. This is essential: prerendered SEO routes must serve their generated HTML rather than the generic SPA shell.")

doc.add_heading("5. Production acceptance checks", level=1)
add_para("Run these checks immediately after deployment and before Search Console submission.")
add_checklist([
    "The homepage loads over HTTPS without redirect loops.",
    "/robots.txt loads and names the correct production sitemap URL.",
    "/sitemap.xml loads and contains 63 URL records.",
    "/resume-examples/software-engineer loads with a unique title and H1.",
    "/student-resume-examples/iit-resume loads with IIT-specific content.",
    "/resume-guides/resume-format-india loads with the full guide content.",
    "Private workspace routes contain noindex,nofollow metadata.",
    "An unknown URL shows the custom not-found page and noindex metadata.",
    "View Source contains the page-specific title, canonical, visible content, and JSON-LD before JavaScript runs.",
    "All canonical URLs use the final production domain and HTTPS.",
])
add_code([
    "curl -I https://YOUR-DOMAIN/robots.txt",
    "curl -I https://YOUR-DOMAIN/sitemap.xml",
    "curl -s https://YOUR-DOMAIN/student-resume-examples/iit-resume",
])

doc.add_heading("6. Configure and validate GA4", level=1)
add_steps([
    "Open GA4 Realtime, then visit the production homepage in a fresh browser tab.",
    "Navigate to a resume example and a resume guide. Confirm page_view events appear.",
    "Click a general resume-start call to action and confirm seo_resume_start appears.",
    "Click Use this template on an SEO detail page and confirm seo_template_use appears.",
    "Inspect event parameters for content_group, content_slug, and template_id.",
    "Mark seo_resume_start and seo_template_use as key events in GA4 if they represent your desired conversions.",
])
add_table(
    ["Event", "Meaning", "Important parameters"],
    [
        ["page_view", "A public or app route was viewed", "page_path, page_title"],
        ["seo_resume_start", "Visitor began from an SEO hub or guide", "content_group, content_slug"],
        ["seo_template_use", "Visitor selected a template from an SEO page", "content_group, content_slug, template_id"],
    ],
    [1.55, 2.65, 2.3],
)

doc.add_heading("7. Set up Google Search Console", level=1)
doc.add_heading("7.1 Add the property", level=2)
add_steps([
    "Open Google Search Console and choose Add property.",
    "Prefer a Domain property when DNS access is available; it covers all protocols and subdomains.",
    "If DNS verification is not available, use a URL-prefix property for the exact HTTPS origin.",
    "Complete the provided DNS TXT or HTML verification method and wait for verification to succeed.",
])
doc.add_heading("7.2 Submit the sitemap", level=2)
add_steps([
    "Select the verified production property.",
    "Open Indexing > Sitemaps.",
    "Enter sitemap.xml and submit it.",
    "Confirm the status changes to Success. Investigate fetch or parsing errors before continuing.",
])
add_callout("Canonical property", "Search Console must match the same HTTPS domain emitted in canonical tags and the sitemap. A www/non-www or subdomain mismatch fragments reporting.")

doc.add_heading("8. Request priority indexing", level=1)
add_para("Use dist/search-console-priority-urls.txt as the launch queue. Search Console limits manual indexing requests, so start with the most commercially important and representative pages.")
add_steps([
    "Inspect the homepage and each main hub first: resume examples, student resumes, templates, and guides.",
    "Inspect IIT, IIM, fresher, internship, ATS template, software engineer, data analyst, and business analyst pages.",
    "Inspect the India resume-format and how-to-make-a-resume guides.",
    "For each URL, run Test Live URL. Resolve crawl or canonical issues before selecting Request Indexing.",
    "Do not submit every URL manually. The sitemap and internal links should handle broad discovery.",
])

doc.add_heading("9. Structured data and search-preview validation", level=1)
add_checklist([
    "Test representative pages with Google Rich Results Test.",
    "Confirm breadcrumb structured data has absolute production URLs.",
    "Confirm FAQ schema appears only where the matching FAQ content is visible.",
    "Confirm guide pages expose Article and BreadcrumbList data.",
    "Check title and description truncation using real search-preview tools or Search Console data.",
    "Confirm Open Graph metadata previews correctly when shared.",
])

doc.add_heading("10. Monitoring schedule", level=1)
add_table(
    ["Timing", "Actions", "Decision rule"],
    [
        ["Launch day", "Verify crawlability, sitemap, GA4, conversions, and core pages", "Fix blocking errors immediately"],
        ["Days 2-7", "Review indexing status and sitemap processing", "Investigate excluded priority URLs"],
        ["Weeks 2-4", "Review queries, pages, CTR, position, and conversions", "Improve pages with impressions but weak engagement"],
        ["Monthly", "Compare 28-day and prior-period performance", "Expand only when data shows a distinct intent"],
        ["Quarterly", "Audit content accuracy, links, templates, schema, and Core Web Vitals", "Refresh or consolidate weak pages"],
    ],
    [1.05, 3.55, 1.9],
)

doc.add_heading("11. Performance metrics to track", level=1)
add_bullets([
    "Indexed URLs versus the 63-URL inventory.",
    "Organic impressions, clicks, click-through rate, and average position by page and query.",
    "Pages ranking in positions 5-20, which are the strongest optimization opportunities.",
    "SEO resume starts and template uses by content group and slug.",
    "Organic visitor-to-resume-start conversion rate.",
    "Search Console Core Web Vitals and mobile usability issues.",
    "Queries receiving impressions without a well-matched landing page.",
])
add_callout("Primary success measure", "Rankings alone are not the goal. Prioritize qualified organic visits that lead to resume starts and template usage.")

doc.add_heading("12. Optimization workflow", level=1)
add_steps([
    "Export Search Console page and query data for the previous 28 days.",
    "Find pages with meaningful impressions and average positions between 5 and 20.",
    "Compare the query intent with the page title, H1, introduction, examples, and call to action.",
    "Improve clarity, specificity, internal links, or examples without keyword stuffing.",
    "Measure for another complete period before making another major change.",
    "Create a new page only when the query represents a genuinely different user need.",
])

doc.add_heading("13. Common problems and responses", level=1)
add_table(
    ["Problem", "Likely cause", "Response"],
    [
        ["Wrong canonical domain", "VITE_SITE_URL missing or incorrect", "Correct environment value and rebuild"],
        ["SEO page shows generic source", "SPA rewrite intercepts static file", "Confirm Vercel filesystem-first routing"],
        ["Sitemap cannot be fetched", "Domain, routing, or robots mismatch", "Open sitemap directly and verify HTTPS URL"],
        ["Discovered, not indexed", "New site, low authority, or similar content", "Strengthen internal links and unique usefulness; wait for recrawl"],
        ["GA4 has no events", "Missing measurement ID or blocked script", "Check production env, Network tab, consent, and Realtime"],
        ["High impressions, low CTR", "Weak or mismatched search snippet", "Rewrite title and description around actual query intent"],
        ["Traffic but no starts", "Page/CTA mismatch or weak trust", "Improve examples, CTA placement, and conversion path"],
    ],
    [1.55, 2.25, 2.7],
)

doc.add_heading("14. Final launch checklist", level=1)
add_checklist([
    "Production domain and GA4 measurement ID are configured.",
    "Working tree was reviewed and only intended files were pushed.",
    "Vercel production build completed successfully.",
    "Robots, sitemap, canonical URLs, prerendered source, and schema were verified.",
    "GA4 page_view, seo_resume_start, and seo_template_use were observed.",
    "Search Console property was verified and sitemap submitted successfully.",
    "Priority URLs were live-tested and indexing requested selectively.",
    "A weekly owner and monthly reporting cadence were assigned.",
])
add_callout("Launch complete", "The SEO release is operational when production pages are crawlable, analytics conversions are visible, the sitemap is accepted, and priority URLs have entered the indexing workflow.", fill=PALE)

doc.core_properties.title = "CareerSense SEO Launch Playbook"
doc.core_properties.subject = "Post-implementation deployment and optimization steps"
doc.core_properties.author = "CareerSense"
doc.core_properties.keywords = "SEO, deployment, Search Console, GA4, Vercel, resume builder"
doc.save(OUT)
print(OUT)
